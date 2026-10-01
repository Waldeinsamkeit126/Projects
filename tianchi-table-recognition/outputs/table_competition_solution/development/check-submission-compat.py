import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
root=Path(__file__).resolve().parent.parent
paths=[root/'submission-candidate-v47-verified-column-counts.xlsx',root/'restart-runs/2026-09-25T03-28-15-719Z-53d112a3/submission-conservative-890-20260925.xlsx']
for path in paths:
    wb=load_workbook(path,read_only=True,data_only=False)
    df=pd.read_excel(path)
    ws=wb.worksheets[0]
    print(json.dumps({'file':path.name,'sheets':wb.sheetnames,'shape':list(df.shape),'headers':list(df.columns),'dtypes':{k:str(v) for k,v in df.dtypes.items()},'nulls':{k:int(v) for k,v in df.isna().sum().items()},'idDuplicates':int(df.iloc[:,0].duplicated().sum()),'firstIdType':ws['A2'].data_type,'lastId':str(df.iloc[-1,0]),'maxAnswerLength':int(df.iloc[:,1].fillna('').astype(str).str.len().max())},ensure_ascii=False))
